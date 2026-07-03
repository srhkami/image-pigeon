import types
import unittest

from unittest.mock import patch

import sys
from pathlib import Path

from typing import Any
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import save_docx


TINY_PNG = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO5wLxQAAAAASUVORK5CYII="


def _sample_request() -> dict[str, Any]:
    return {
        "title": "story6",
        "font_size": "12",
        "align_vertical": "center",
        "mode": "4",
        "images": [
            {
                "id": "item-1",
                "base64": TINY_PNG,
                "remark": "一",
                "rotation": 0,
            }
        ],
        "path": "/tmp/story6.docx",
    }


def _slot(item_id: str | None, role: str, order: int) -> dict[str, Any]:
    return {
        "slotId": f"slot-{order}",
        "itemId": item_id,
        "role": role,
        "order": order,
    }


def _page(template: str, slots: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "pageId": f"page-{template}",
        "template": template,
        "slots": slots,
    }


TEMPLATE_SLOTS = {
    "landscape-2": [
        _slot("item-1", "landscape-top", 1),
        _slot(None, "landscape-bottom", 2),
    ],
    "portrait-large-2": [
        _slot("item-1", "portrait-large-left", 1),
        _slot(None, "portrait-large-right", 2),
    ],
    "portrait-small-6": [
        _slot("item-1", "portrait-small-1", 1),
        _slot(None, "portrait-small-2", 2),
        _slot(None, "portrait-small-3", 3),
        _slot(None, "portrait-small-4", 4),
        _slot(None, "portrait-small-5", 5),
        _slot(None, "portrait-small-6", 6),
    ],
    "mixed-landscape1-small3": [
        _slot("item-1", "mixed-landscape-top", 1),
        _slot(None, "mixed-small-bottom-left", 2),
        _slot(None, "mixed-small-bottom-middle", 3),
        _slot(None, "mixed-small-bottom-right", 4),
    ],
    "mixed-small3-landscape1": [
        _slot(None, "mixed-small-top-left", 1),
        _slot(None, "mixed-small-top-middle", 2),
        _slot(None, "mixed-small-top-right", 3),
        _slot("item-1", "mixed-landscape-bottom", 4),
    ],
}


class SaveDocxCollageTest(unittest.TestCase):
    def _auto_request(self, pages: list[dict[str, Any]] | None = None) -> dict[str, Any]:
        request = _sample_request()
        request["layoutMode"] = "auto-collage-v1"
        request["pages"] = pages if pages is not None else [
            _page("landscape-2", TEMPLATE_SLOTS["landscape-2"]),
        ]
        return request

    def test_output_word_auto_payload_parses_pages(self):
        data = save_docx.OutputWord(self._auto_request())

        self.assertEqual(data.layout_mode, "auto-collage-v1")
        self.assertEqual(len(data.pages), 1)
        self.assertEqual(data.mode, 4)

    def test_output_word_legacy_mode_fallback_still_parses(self):
        data = save_docx.OutputWord(_sample_request())

        self.assertIsNone(data.layout_mode)
        self.assertEqual(data.pages, [])
        self.assertEqual(data.mode, 4)

    def test_creat_docx_uses_auto_renderer_when_auto_layout_present(self):
        with patch.object(save_docx.OutputWord, "to_compressed_images", return_value=[]), \
             patch.object(save_docx, "_render_auto_collage_pages", return_value="auto") as render_mock:
            data = save_docx.OutputWord(self._auto_request())
            doc = save_docx.creat_docx(data)

        render_mock.assert_called_once()
        self.assertEqual(doc, "auto")

    def test_creat_docx_falls_back_to_legacy_when_auto_pages_empty(self):
        with patch.object(save_docx.OutputWord, "to_compressed_images", return_value=[]), \
             patch.object(save_docx, "_build_legacy_layout", return_value="legacy") as legacy_mock:
            data = save_docx.OutputWord(self._auto_request([]))
            doc = save_docx.creat_docx(data)

        legacy_mock.assert_called_once()
        self.assertEqual(doc, "legacy")

    def test_slot_image_map_resolves_by_item_id(self):
        image_map = save_docx._build_image_map(
            [
                types.SimpleNamespace(id="item-1", value=1),
                types.SimpleNamespace(id="item-2", value=2),
            ]
        )
        image = save_docx._resolve_slot_image(image_map, {"itemId": "item-2"})

        self.assertIsNotNone(image)
        self.assertEqual(image.value, 2)

    def test_each_auto_template_renders_one_table_without_image_io(self):
        for template, slots in TEMPLATE_SLOTS.items():
            with self.subTest(template=template):
                request = self._auto_request([_page(template, slots)])
                images = [types.SimpleNamespace(id="item-1", remark="測試")]

                with patch.object(save_docx.OutputWord, "to_compressed_images", return_value=images), \
                     patch.object(save_docx, "handle_table_write") as write_mock:
                    data = save_docx.OutputWord(request)
                    doc = save_docx.creat_docx(data)

                self.assertEqual(len(doc.tables), 1)
                write_mock.assert_called_once()

    def test_empty_slots_keep_table_and_skip_image_write(self):
        request = self._auto_request([
            _page("landscape-2", [
                _slot(None, "landscape-top", 1),
                _slot(None, "landscape-bottom", 2),
            ])
        ])

        with patch.object(save_docx.OutputWord, "to_compressed_images", return_value=[]), \
             patch.object(save_docx, "handle_table_write") as write_mock:
            data = save_docx.OutputWord(request)
            doc = save_docx.creat_docx(data)

        self.assertEqual(len(doc.tables), 1)
        write_mock.assert_not_called()

    def test_portrait_small_rows_fit_two_rows_on_one_page(self):
        doc = Document()

        save_docx._append_portrait_small_6_table(
            doc,
            align="center",
            images_by_id={},
            slots=TEMPLATE_SLOTS["portrait-small-6"],
            image_index=1,
        )

        table = doc.tables[0]
        total_height = sum(int(row.height) for row in table.rows if row.height is not None)
        self.assertLessEqual(total_height, int(Cm(23)))
        actual_heights = [int(row.height) for row in table.rows if row.height is not None]
        expected_heights = [int(Cm(value)) for value in [7.2, 0.6, 3.2, 7.2, 0.6, 3.2]]
        for actual, expected in zip(actual_heights, expected_heights, strict=True):
            self.assertAlmostEqual(actual, expected, delta=500)

    def test_portrait_small_columns_are_evenly_distributed(self):
        renderers = [
            (save_docx._append_portrait_small_6_table, TEMPLATE_SLOTS["portrait-small-6"]),
            (save_docx._append_mixed_landscape1_small3_table, TEMPLATE_SLOTS["mixed-landscape1-small3"]),
            (save_docx._append_mixed_small3_landscape1_table, TEMPLATE_SLOTS["mixed-small3-landscape1"]),
        ]

        for renderer, slots in renderers:
            with self.subTest(renderer=renderer.__name__):
                doc = Document()
                renderer(
                    doc,
                    align="center",
                    images_by_id={},
                    slots=slots,
                    image_index=1,
                )

                table = doc.tables[0]
                grid_widths = [int(col.get(qn('w:w'))) for col in table._tbl.tblGrid.gridCol_lst]
                expected_width = int(Cm(5.2).twips)
                self.assertEqual(len(grid_widths), 3)
                for grid_width in grid_widths:
                    self.assertAlmostEqual(grid_width, expected_width, delta=1)
                self.assertEqual(len(set(grid_widths)), 1)


    def test_mixed_portrait_small_and_landscape_rows_are_role_aware(self):
        doc = Document()
        save_docx._append_mixed_landscape1_small3_table(
            doc,
            align="center",
            images_by_id={},
            slots=TEMPLATE_SLOTS["mixed-landscape1-small3"],
            image_index=1,
        )
        top_landscape_rows = [int(row.height) for row in doc.tables[0].rows if row.height is not None]
        self.assertEqual(len(doc.tables[0].rows), 5)
        expected_top_landscape = [int(Cm(value)) for value in [8.2, 3.2, 7.6, 0.6, 3.2]]
        for actual, expected in zip(top_landscape_rows, expected_top_landscape, strict=True):
            self.assertAlmostEqual(actual, expected, delta=500)
        self.assertAlmostEqual(sum(top_landscape_rows[:2]), sum(top_landscape_rows[2:]), delta=1000)

        doc = Document()
        save_docx._append_mixed_small3_landscape1_table(
            doc,
            align="center",
            images_by_id={},
            slots=TEMPLATE_SLOTS["mixed-small3-landscape1"],
            image_index=1,
        )
        bottom_landscape_rows = [int(row.height) for row in doc.tables[0].rows if row.height is not None]
        self.assertEqual(len(doc.tables[0].rows), 5)
        expected_bottom_landscape = [int(Cm(value)) for value in [7.6, 0.6, 3.2, 8.2, 3.2]]
        for actual, expected in zip(bottom_landscape_rows, expected_bottom_landscape, strict=True):
            self.assertAlmostEqual(actual, expected, delta=500)
        self.assertAlmostEqual(sum(bottom_landscape_rows[:3]), sum(bottom_landscape_rows[3:]), delta=1000)

    def test_landscape_number_column_keeps_minimum_width(self):
        doc = Document()

        save_docx._append_landscape_2_table(
            doc,
            align="center",
            images_by_id={},
            slots=TEMPLATE_SLOTS["landscape-2"],
            image_index=1,
        )

        table = doc.tables[0]
        grid_widths = [int(col.get(qn('w:w'))) for col in table._tbl.tblGrid.gridCol_lst]
        number_width = table.cell(1, 0).width
        remark_width = table.cell(1, 1).width
        assert number_width is not None
        assert remark_width is not None
        self.assertAlmostEqual(grid_widths[0], int(Cm(1.8).twips), delta=1)
        self.assertAlmostEqual(int(number_width), int(Cm(1.8)), delta=500)
        self.assertGreater(grid_widths[1], grid_widths[0])
        self.assertGreater(int(remark_width), int(number_width))

    def test_auto_numbering_continues_across_pages(self):
        request = self._auto_request([
            _page("landscape-2", [
                _slot("item-1", "landscape-top", 1),
                _slot("item-2", "landscape-bottom", 2),
            ]),
            _page("portrait-large-2", [
                _slot("item-3", "portrait-large-left", 1),
                _slot(None, "portrait-large-right", 2),
            ]),
        ])
        images = [
            types.SimpleNamespace(id="item-1", remark="一"),
            types.SimpleNamespace(id="item-2", remark="二"),
            types.SimpleNamespace(id="item-3", remark="三"),
        ]
        seen_indexes: list[int] = []

        def fake_write(**kwargs):
            seen_indexes.append(kwargs["index"])

        with patch.object(save_docx.OutputWord, "to_compressed_images", return_value=images), \
             patch.object(save_docx, "handle_table_write", side_effect=fake_write):
            data = save_docx.OutputWord(request)
            doc = save_docx.creat_docx(data)

        self.assertEqual(len(doc.tables), 2)
        self.assertEqual(seen_indexes, [1, 2, 3])


if __name__ == "__main__":
    unittest.main()
